"""Builds anbit_leads.xlsx from leads.tsv (Instagram prospects for Anbit)."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
SRC = HERE / "leads.tsv"
OUT = HERE / "anbit_leads.xlsx"
COLLECTED_ON = "08/10/2026"

FONT = "Arial"
BRAND = "1F3A5F"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor=BRAND)
INPUT_HEAD_FILL = PatternFill("solid", fgColor="BF8F00")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUSES = [
    "Δεν στάλθηκε",
    "Στάλθηκε",
    "Follow-up στάλθηκε",
    "Ενδιαφέρεται",
    "Ραντεβού / Demo",
    "Συνεργάτης",
    "Δεν ενδιαφέρεται",
    "Δεν απάντησε",
]

# Priority A: venues with table service and frequent repeat visits (cafés, all-day, brunch).
# Priority C: mostly take-away, pastry shops, or chains/franchises with their own systems.
# Priority B: everything else with table service (bars, tavernas, meze, restaurants, street food).
A_KEYS = ("cafe", "café", "coffee", "brunch", "all-day", "καφέ", "καφενείο", "καφετέρια", "bistro", "bakery")
C_KEYS = ("take away", "ζαχαροπλαστείο", "αλυσίδα", "franchise")


def priority(category, info):
    text = f"{category} {info}".lower()
    if any(k in text for k in C_KEYS):
        return "C"
    if any(k in category.lower() for k in A_KEYS):
        return "A"
    return "B"


def load_rows():
    with SRC.open(encoding="utf-8") as f:
        reader = csv.reader(f, delimiter="\t")
        next(reader)
        rows = []
        for r in reader:
            handle, name, city, area, category, info, verify = (r + [""] * 7)[:7]
            rows.append(
                dict(handle=handle, name=name, city=city, area=area, category=category,
                     info=info, verify=verify, prio=priority(category, info))
            )
    city_order = {"Αθήνα": 0, "Θεσσαλονίκη": 1}
    rows.sort(key=lambda r: (city_order.get(r["city"], 9), r["prio"], r["name"].lower()))
    return rows


def style_header(cell, fill=HEAD_FILL):
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def build_messages(wb):
    ws = wb.create_sheet("Μηνύματα")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110

    ws["A1"] = "Πρότυπα μηνυμάτων για Instagram DM"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=BRAND)
    ws["A2"] = ("Άλλαξε το κείμενο εδώ και ενημερώνεται αυτόματα η στήλη «Έτοιμο μήνυμα» στο φύλλο Λογαριασμοί. "
                "Το {ΚΑΤΑΣΤΗΜΑ} γίνεται το όνομα του μαγαζιού και το {ΟΝΟΜΑ} το όνομά σου.")
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")
    ws.merge_cells("A2:B2")

    ws["A4"] = "Το όνομά σου (συμπλήρωσε)"
    ws["B4"] = "[Το όνομά σου]"
    ws["B4"].fill = INPUT_FILL

    templates = [
        ("1ο μήνυμα – Καφέ / All-day / Brunch (προτεραιότητα A)",
         "Γεια σας! 👋 Είμαι ο/η {ΟΝΟΜΑ} από την Anbit (anbit.gr). Μας άρεσε πολύ το {ΚΑΤΑΣΤΗΜΑ} και θα θέλαμε να σας προτείνουμε μια συνεργασία.\n\n"
         "Η Anbit είναι παραγγελία από το τραπέζι με QR + σύστημα επιβράβευσης πελατών:\n"
         "• Ο πελάτης σκανάρει το QR, βλέπει το μενού με φωτογραφίες (Ελληνικά/Αγγλικά) και παραγγέλνει, χωρίς εφαρμογή\n"
         "• Η παραγγελία έρχεται στο κινητό ή στο tablet σας με τον αριθμό τραπεζιού\n"
         "• Με κάθε παραγγελία μαζεύει πόντους (XP) που εξαργυρώνονται μόνο στο δικό σας μαγαζί, οπότε έχει λόγο να ξαναέρθει\n"
         "• Βλέπετε ποιος έρχεται, πόσο συχνά και τι παίρνει\n"
         "• Δεν αλλάζει τίποτα στο ταμείο: ίδιο POS, 0% προμήθεια, το στήνουμε εμείς σε 10 λεπτά\n"
         "• Δωρεάν προβολή στα TikTok/Instagram της Anbit και στον χάρτη της Anbit\n\n"
         "Θα σας ενδιέφερε να σας το δείξουμε από κοντά σε 10 λεπτά, όποτε σας βολεύει;"),
        ("1ο μήνυμα – Bar / Ταβέρνα / Εστιατόριο / Street food (προτεραιότητα B, C)",
         "Γεια σας! 👋 Είμαι ο/η {ΟΝΟΜΑ} από την Anbit (anbit.gr). Είδαμε το {ΚΑΤΑΣΤΗΜΑ} και θα θέλαμε να σας προτείνουμε μια συνεργασία.\n\n"
         "Με την Anbit ο πελάτης σκανάρει ένα QR στο τραπέζι, βλέπει το μενού με φωτογραφίες σε Ελληνικά/Αγγλικά και παραγγέλνει χωρίς εφαρμογή. "
         "Η παραγγελία έρχεται κατευθείαν σε εσάς με τον αριθμό τραπεζιού, οπότε το προσωπικό κερδίζει χρόνο τις ώρες αιχμής.\n\n"
         "Επιπλέον ο πελάτης κερδίζει πόντους που εξαργυρώνει μόνο σε εσάς, οπότε έχει λόγο να επιστρέψει. "
         "Ίδιο POS, ίδια απόδειξη, 0% προμήθεια, και το στήνουμε εμείς σε 10 λεπτά. Σας προβάλλουμε επίσης δωρεάν στα social της Anbit.\n\n"
         "Θα θέλατε να σας στείλουμε περισσότερα ή να περάσουμε για μια σύντομη παρουσίαση;"),
        ("Follow-up (3-4 μέρες μετά, αν δεν απαντήσουν)",
         "Καλησπέρα ξανά από την Anbit! 😊 Ήθελα απλώς να δω αν είδατε το μήνυμά μου για το QR ordering με επιβράβευση πελατών για το {ΚΑΤΑΣΤΗΜΑ}. "
         "Αν σας ενδιαφέρει, περνάμε όποτε σας βολεύει για 10 λεπτά, χωρίς καμία δέσμευση. Μπορείτε να δείτε και περισσότερα στο anbit.gr"),
        ("Απάντηση: «Πόσο κοστίζει;»",
         "Ευχαριστούμε για το ενδιαφέρον! Η Anbit δεν παίρνει καμία προμήθεια από τις παραγγελίες (0%), αφού ο πελάτης πληρώνει όπως πάντα στο ταμείο σας. "
         "Το κόστος είναι [ΣΥΜΠΛΗΡΩΣΕ ΤΙΜΗ/ΠΑΚΕΤΟ]. Θέλετε να κλείσουμε μια σύντομη παρουσίαση για να το δείτε στην πράξη στο {ΚΑΤΑΣΤΗΜΑ};"),
    ]
    ws["A6"] = "Πρότυπο"
    ws["B6"] = "Κείμενο"
    style_header(ws["A6"])
    style_header(ws["B6"])
    for i, (label, text) in enumerate(templates, start=7):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=text).fill = INPUT_FILL
        ws.row_dimensions[i].height = 230 if i < 9 else 75
    for row in ws.iter_rows(min_row=4, max_row=6 + len(templates)):
        for c in row:
            if c.row != 6:
                c.font = Font(name=FONT, bold=(c.column == 1))
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def build_leads(wb, rows):
    ws = wb.active
    ws.title = "Λογαριασμοί"
    headers = [
        ("#", 5), ("Κατάστημα", 30), ("Instagram", 26), ("Προφίλ", 12), ("Στείλε DM", 12),
        ("Πόλη", 13), ("Περιοχή / Διεύθυνση", 32), ("Κατηγορία", 26), ("Προτεραιότητα", 13),
        ("Στοιχεία από το bio", 42), ("Προς επιβεβαίωση", 20), ("Κατάσταση", 20),
        ("Ημ/νία 1ου μηνύματος", 14), ("Follow-up στις", 14), ("Σημειώσεις", 34), ("Έτοιμο μήνυμα", 60),
    ]
    input_cols = {"Κατάσταση", "Ημ/νία 1ου μηνύματος", "Σημειώσεις"}
    HR = 4  # header row
    first, last = HR + 1, HR + len(rows)

    ws["A1"] = "Anbit – Λογαριασμοί Instagram για συνεργασία (Αθήνα & Θεσσαλονίκη)"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=BRAND)
    ws["A2"] = ("Συμπληρώνεις μόνο τις κίτρινες στήλες: Κατάσταση, Ημ/νία 1ου μηνύματος, Σημειώσεις. "
                "Το «Στείλε DM» ανοίγει κατευθείαν συνομιλία στο Instagram. Το «Έτοιμο μήνυμα» βγαίνει από το φύλλο Μηνύματα.")
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")

    for col, (h, w) in enumerate(headers, start=1):
        c = ws.cell(row=HR, column=col, value=h)
        style_header(c, INPUT_HEAD_FILL if h in input_cols else HEAD_FILL)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[HR].height = 32

    link_font = Font(name=FONT, color="0563C1", underline="single")
    for i, r in enumerate(rows):
        n = first + i
        values = [
            i + 1, r["name"], "@" + r["handle"], "Άνοιγμα", "Μήνυμα",
            r["city"], r["area"], r["category"], r["prio"], r["info"], r["verify"],
            "Δεν στάλθηκε", None, f'=IF(M{n}="","",M{n}+4)', None,
            (f"=SUBSTITUTE(SUBSTITUTE(IF(I{n}=\"A\",'Μηνύματα'!$B$7,'Μηνύματα'!$B$8),"
             f"\"{{ΚΑΤΑΣΤΗΜΑ}}\",B{n}),\"{{ΟΝΟΜΑ}}\",'Μηνύματα'!$B$4)"),
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=n, column=col, value=v)
            c.font = Font(name=FONT)
            c.border = BORDER
            c.alignment = Alignment(vertical="top")
        ws.cell(row=n, column=4).hyperlink = f"https://www.instagram.com/{r['handle']}/"
        ws.cell(row=n, column=4).font = link_font
        ws.cell(row=n, column=5).hyperlink = f"https://ig.me/m/{r['handle']}"
        ws.cell(row=n, column=5).font = link_font
        ws.cell(row=n, column=9).alignment = Alignment(horizontal="center", vertical="top")
        for col in (12, 13, 15):
            ws.cell(row=n, column=col).fill = INPUT_FILL
        for col in (13, 14):
            ws.cell(row=n, column=col).number_format = "DD/MM/YYYY"
        if r["verify"]:
            ws.cell(row=n, column=11).font = Font(name=FONT, color="C00000")

    status_dv = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=True)
    status_dv.add(f"L{first}:L{last}")
    prio_dv = DataValidation(type="list", formula1='"A,B,C"', allow_blank=False)
    prio_dv.add(f"I{first}:I{last}")
    date_dv = DataValidation(type="date", operator="greaterThan", formula1="45000", allow_blank=True)
    date_dv.add(f"M{first}:M{last}")
    for dv in (status_dv, prio_dv, date_dv):
        ws.add_data_validation(dv)

    rng = f"A{first}:P{last}"
    for status, color in [("Ενδιαφέρεται", "C6EFCE"), ("Ραντεβού / Demo", "A9D08E"), ("Συνεργάτης", "70AD47"),
                          ("Δεν ενδιαφέρεται", "F4CCCC"), ("Στάλθηκε", "DDEBF7"), ("Follow-up στάλθηκε", "BDD7EE")]:
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'$L{first}="{status}"'], fill=PatternFill("solid", fgColor=color)))
    prio_colors = {"A": "C6EFCE", "B": "FFEB9C", "C": "EDEDED"}
    for p, color in prio_colors.items():
        ws.conditional_formatting.add(
            f"I{first}:I{last}",
            FormulaRule(formula=[f'$I{first}="{p}"'], fill=PatternFill("solid", fgColor=color)))

    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{HR}:P{last}"
    return first, last


def build_summary(wb, first, last):
    ws = wb.create_sheet("Σύνοψη", 0)
    ref = lambda col: f"'Λογαριασμοί'!${col}${first}:${col}${last}"
    ws.column_dimensions["A"].width = 26
    for col in "BCDE":
        ws.column_dimensions[col].width = 14

    ws["A1"] = "Anbit – Σύνοψη καμπάνιας Instagram"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=BRAND)
    ws["A2"] = f"Λογαριασμοί που συγκεντρώθηκαν στις {COLLECTED_ON}. Οι αριθμοί ενημερώνονται μόνοι τους."
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")

    # City x priority
    for col, h in enumerate(["Πόλη", "A", "B", "C", "Σύνολο"], start=1):
        style_header(ws.cell(row=4, column=col, value=h))
    cities = ["Αθήνα", "Θεσσαλονίκη"]
    for i, city in enumerate(cities, start=5):
        ws.cell(row=i, column=1, value=city)
        for j, p in enumerate("ABC", start=2):
            ws.cell(row=i, column=j, value=f'=COUNTIFS({ref("F")},$A{i},{ref("I")},"{p}")')
        ws.cell(row=i, column=5, value=f"=SUM(B{i}:D{i})")
    tr = 5 + len(cities)
    ws.cell(row=tr, column=1, value="Σύνολο")
    for j in range(2, 6):
        L = get_column_letter(j)
        ws.cell(row=tr, column=j, value=f"=SUM({L}5:{L}{tr - 1})")

    # Status pipeline
    sr = tr + 2
    for col, h in enumerate(["Κατάσταση", "Αθήνα", "Θεσσαλονίκη", "Σύνολο"], start=1):
        style_header(ws.cell(row=sr, column=col, value=h))
    for i, s in enumerate(STATUSES, start=sr + 1):
        ws.cell(row=i, column=1, value=s)
        ws.cell(row=i, column=2, value=f'=COUNTIFS({ref("L")},$A{i},{ref("F")},"Αθήνα")')
        ws.cell(row=i, column=3, value=f'=COUNTIFS({ref("L")},$A{i},{ref("F")},"Θεσσαλονίκη")')
        ws.cell(row=i, column=4, value=f"=B{i}+C{i}")
    end = sr + len(STATUSES)
    rr = end + 2
    ws.cell(row=rr, column=1, value="Ποσοστό απάντησης")
    sent = f'(COUNTA({ref("L")})-COUNTIF({ref("L")},"Δεν στάλθηκε"))'
    replied = "+".join(f'COUNTIF({ref("L")},"{s}")' for s in
                       ["Ενδιαφέρεται", "Ραντεβού / Demo", "Συνεργάτης", "Δεν ενδιαφέρεται"])
    ws.cell(row=rr, column=2, value=f"=IF({sent}=0,0,({replied})/{sent})").number_format = "0.0%"
    ws.cell(row=rr, column=3, value="απαντήσεις ÷ μηνύματα που στάλθηκαν")

    for row in ws.iter_rows(min_row=5, max_row=rr):
        for c in row:
            if c.row in (sr,):
                continue
            c.font = Font(name=FONT, bold=(c.row == tr or c.column == 1))
            if c.row != rr and c.value is not None:
                c.border = BORDER
    ws.cell(row=rr, column=3).font = Font(name=FONT, italic=True, color="595959")

    # Guide
    gr = rr + 2
    ws.cell(row=gr, column=1, value="Πώς βγήκε η λίστα & πώς τη χρησιμοποιείς").font = Font(
        name=FONT, bold=True, size=12, color=BRAND)
    notes = [
        "• Δημόσια επαγγελματικά προφίλ Instagram (καφέ, all-day, brunch, bars, μεζεδοπωλεία, ταβέρνες, street food) σε Αθήνα και Θεσσαλονίκη, "
        "μαζί με όσα στοιχεία γράφουν τα ίδια στο bio τους (διεύθυνση, ωράριο, τηλέφωνο).",
        "• Προτεραιότητα A: καφέ / all-day / brunch με τραπέζια και πελάτες που έρχονται συχνά, άρα το καλύτερο ταίριασμα για QR + πόντους. "
        "B: bars, ταβέρνες, μεζεδοπωλεία, εστιατόρια, street food. C: κυρίως take away, ζαχαροπλαστεία ή αλυσίδες/franchise.",
        "• Η στήλη «Προς επιβεβαίωση» σημειώνει όσα δεν έγραφαν καθαρά περιοχή ή είδος μαγαζιού. Ρίξε μια ματιά στο προφίλ πριν στείλεις.",
        "• Στείλε από τον επαγγελματικό λογαριασμό της Anbit. Τα μηνύματα σε λογαριασμούς που δεν σε ακολουθούν πάνε στα «Αιτήματα μηνυμάτων», "
        "οπότε ένα follow και ένα like πριν το μήνυμα βοηθάνε να το δουν.",
        "• Μην στέλνεις μαζικά: περίπου 20-30 νέα μηνύματα την ημέρα, με λίγο διαφορετικό κείμενο, για να μη σε περιορίσει το Instagram.",
        "• Γράψε την ημερομηνία στο «Ημ/νία 1ου μηνύματος» και το «Follow-up στις» βγαίνει μόνο του 4 μέρες μετά.",
    ]
    for k, t in enumerate(notes, start=gr + 1):
        c = ws.cell(row=k, column=1, value=t)
        c.font = Font(name=FONT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=k, start_column=1, end_row=k, end_column=5)
        ws.row_dimensions[k].height = 48


def main():
    rows = load_rows()
    wb = Workbook()
    first, last = build_leads(wb, rows)
    build_messages(wb)
    build_summary(wb, first, last)
    wb.active = 0
    wb.save(OUT)
    print(f"{len(rows)} rows -> {OUT}")


if __name__ == "__main__":
    main()
